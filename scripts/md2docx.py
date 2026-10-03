# md2docx.py — simple Markdown -> Word converter for the action report.
# Usage: python md2docx.py <input.md> <output.docx>
import io
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

src, dst = sys.argv[1], sys.argv[2]
md = io.open(src, encoding="utf-8").read()

doc = Document()
st = doc.styles["Normal"]
st.font.name = "微软雅黑"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

CLEAN = re.compile(r"\*\*|`")


def clean(s):
    return CLEAN.sub("", s).strip()


def add_table(rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            if i < len(cells):
                cells[i].text = clean(v)


lines = md.split("\n")
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    # table block
    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                rows.append(cells)
            i += 1
        if rows:
            add_table(rows)
        doc.add_paragraph()
        continue

    if ln.startswith("# "):
        doc.add_heading(clean(ln[2:]), level=0)
    elif ln.startswith("## "):
        doc.add_heading(clean(ln[3:]), level=1)
    elif ln.startswith("### "):
        doc.add_heading(clean(ln[4:]), level=2)
    elif ln.startswith("> "):
        p = doc.add_paragraph()
        r = p.add_run(clean(ln[2:]))
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    elif re.match(r"^[-*] ", ln):
        p = doc.add_paragraph(style="List Bullet")
        txt = clean(ln[2:])
        m = re.match(r"^(⚠️|✅|🔑|💡)?\s*(.*)$", txt)
        p.add_run(txt)
    elif re.match(r"^\d+\. ", ln):
        doc.add_paragraph(clean(re.sub(r"^\d+\. ", "", ln)), style="List Number")
    elif ln.strip() == "---":
        pass
    elif ln.strip():
        doc.add_paragraph(clean(ln))
    i += 1

doc.save(dst)
print(f"saved {dst} | paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
